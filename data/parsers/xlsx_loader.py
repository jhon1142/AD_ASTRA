"""
Parser para archivos Excel (.xlsx) — AD_ASTRA CODEFEST 2026.

Reglas:
- El archivo completo representa un único documento.
- Todas las filas comparten el mismo doc_id.
- Cada fila se conserva como unidad independiente de fragmentación.
- Cada valor se representa como columna: valor.
"""

from __future__ import annotations

from pathlib import Path
from typing import Union

from core.document import Document
from data.parsers.base_loader import (
    BaseLoader,
    infer_fenomeno,
    make_doc_id,
)


class XLSXLoader(BaseLoader):
    """
    Parsea archivos Excel preservando la identidad documental.

    Todas las filas de un mismo archivo comparten el mismo doc_id.
    SentenceSplitter se encarga posteriormente de convertirlas en
    chunks independientes.
    """

    def __init__(
        self,
        sheet_name: str | int | None = None,
        content_columns: list[str] | None = None,
        header_row: int = 0,
        skip_empty: bool = True,
    ) -> None:

        self.sheet_name = sheet_name
        self.content_columns = content_columns
        self.header_row = header_row
        self.skip_empty = skip_empty

    def load(
        self,
        source: Union[str, Path],
    ) -> list[Document]:

        try:
            import openpyxl

        except ImportError as exc:

            raise ImportError(
                "Instala 'openpyxl': "
                "pip install openpyxl"
            ) from exc

        path = Path(source)

        # Un único identificador por archivo.
        doc_id = make_doc_id(
            path
        )

        fenomeno = infer_fenomeno(
            path
        )

        workbook = openpyxl.load_workbook(
            str(path),
            read_only=True,
            data_only=True,
        )

        if isinstance(
            self.sheet_name,
            int,
        ):

            sheet = workbook.worksheets[
                self.sheet_name
            ]

        elif self.sheet_name is None:

            sheet = workbook.active

        else:

            sheet = workbook[
                self.sheet_name
            ]

        sheet_title = sheet.title

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        workbook.close()

        if not rows:
            return []

        if self.header_row >= len(rows):
            raise ValueError(
                f"header_row={self.header_row} "
                "está fuera del rango de la hoja."
            )

        headers = [
            (
                str(header)
                if header is not None
                else f"col_{index}"
            )
            for index, header
            in enumerate(
                rows[self.header_row]
            )
        ]

        data_rows = rows[
            self.header_row + 1:
        ]

        columns = (
            self.content_columns
            or headers
        )

        documents: list[Document] = []

        for row_num, row in enumerate(
            data_rows,
            start=1,
        ):

            row_dict = dict(
                zip(
                    headers,
                    row,
                )
            )

            parts: list[str] = []

            for column in columns:

                value = row_dict.get(
                    column
                )

                if value is None:
                    continue

                value_str = str(
                    value
                ).strip()

                if not value_str:
                    continue

                parts.append(
                    f"{column}: {value_str}"
                )

            if (
                self.skip_empty
                and not parts
            ):
                continue

            content = " | ".join(
                parts
            )

            # Fila física en Excel.
            excel_row = (
                self.header_row
                + 1
                + row_num
            )

            documents.append(
                Document(
                    # MISMO doc_id para todas las filas.
                    doc_id=doc_id,

                    fuente=path.name,
                    formato="xlsx",
                    fenomeno=fenomeno,

                    content=content,

                    metadata={
                        "ruta_completa": str(path),
                        "hoja": sheet_title,

                        # Número real de fila en Excel.
                        "fila": excel_row,

                        # Índice de fila de datos base 0.
                        "fila_indice": row_num - 1,

                        "tipo_unidad": "fila",
                    },
                )
            )

        return documents